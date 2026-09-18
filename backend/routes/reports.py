# backend/routes/reports.py
from flask import Blueprint, request, send_file, jsonify
from io import BytesIO
from datetime import datetime, date as date_cls
import calendar
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import or_, and_

from models import db, DentalEntry, Doctor, Hospital, BalanceCarry
from utils.jwt_utils import token_required

report_bp = Blueprint('reports', __name__, url_prefix='/api/reports')


def _format_entry_period(entry):
    et = getattr(entry, 'entry_type', 'date') or 'date'
    em = getattr(entry, 'entry_month', None)
    if et == 'month' and em and len(str(em)) == 7:
        try:
            y, m = str(em).split('-')
            d = date_cls(int(y), int(m), 1)
            return d.strftime('%B %Y')
        except (ValueError, TypeError):
            pass
    return entry.entry_date.isoformat() if entry.entry_date else ''


@report_bp.route('/excel', methods=['POST'])
@token_required
def export_excel():
    try:
        data = request.get_json(silent=True) or {}

        doctor_id = data.get('doctor_id')
        hospital_id = data.get('hospital_id')
        date_from = data.get('date_from')
        date_to = data.get('date_to')
        month = data.get('month')

        if not doctor_id:
            return jsonify({'error': 'doctor_id is required'}), 400
        if not month and (not date_from or not date_to):
            return jsonify({'error': 'Provide either month or date_from/date_to'}), 400

        # ---- Build query (DATE-only entries) ----
        query = DentalEntry.query.filter(DentalEntry.doctor_id == int(doctor_id))
        if hospital_id:
            query = query.filter(DentalEntry.hospital_id == int(hospital_id))

        query = query.filter(DentalEntry.entry_type == 'date')

        period_text = ''
        filename_suffix = ''

        if month:
            try:
                y_s, m_s = str(month).split('-')
                year_i, mon_i = int(y_s), int(m_s)
                first_day = date_cls(year_i, mon_i, 1)
                last_day = date_cls(year_i, mon_i, calendar.monthrange(year_i, mon_i)[1])
            except (ValueError, TypeError):
                return jsonify({'error': "month must be in 'YYYY-MM' format"}), 400

            query = query.filter(
                db.func.date(DentalEntry.entry_date) >= first_day
            ).filter(
                db.func.date(DentalEntry.entry_date) <= last_day
            )
            period_text = first_day.strftime('%B %Y')
            filename_suffix = str(month)

        else:
            try:
                date_from_obj = datetime.strptime(date_from, '%Y-%m-%d').date()
                date_to_obj = datetime.strptime(date_to, '%Y-%m-%d').date()
            except ValueError:
                return jsonify({'error': 'Dates must be YYYY-MM-DD'}), 400

            query = query.filter(db.func.date(DentalEntry.entry_date) >= date_from_obj)
            query = query.filter(db.func.date(DentalEntry.entry_date) <= date_to_obj)
            period_text = f'{date_from} to {date_to}'
            filename_suffix = f'{date_from}_to_{date_to}'

        query = query.order_by(DentalEntry.entry_date.desc())
        entries = query.all()

        # ---- Look up names ----
        doctor = Doctor.query.get(int(doctor_id))
        doctor_name = doctor.doctor_name if doctor else ''

        hospital_name = 'All'
        if hospital_id:
            hospital = Hospital.query.get(int(hospital_id))
            if hospital:
                hospital_name = hospital.hospital_name

        # ---- Fetch previous balance for the "Remaining" calc ----
        prev_balance = 0.0
        if month:
            q = BalanceCarry.query.filter_by(doctor_id=int(doctor_id), month=str(month))
            if hospital_id:
                q = q.filter_by(hospital_id=int(hospital_id))
            else:
                q = q.filter(BalanceCarry.hospital_id.is_(None))
            row = q.first()
            if row:
                prev_balance = float(row.previous_balance or 0)

        # ---- Build workbook ----
        wb = Workbook()
        ws = wb.active
        ws.title = 'Revenue'

        title_font = Font(bold=True, size=14, color='1F3A3D')
        sub_font = Font(size=10, color='555555')
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill('solid', fgColor='1F3A3D')
        total_fill = PatternFill('solid', fgColor='E8E3D6')
        thin = Side(style='thin', color='BBBBBB')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal='center', vertical='center')
        right = Alignment(horizontal='right', vertical='center')

        # Title block
        ws.merge_cells('A1:G1')
        ws['A1'] = 'THE DENTAL ART LABORATORY'
        ws['A1'].font = title_font
        ws['A1'].alignment = center

        ws.merge_cells('A2:G2')
        ws['A2'] = ('Kamala Enclave, 3rd Floor, Near Kanyakha Homes, '
                    'Kugler Hospital Road, Kothapet, GUNTUR-522001.')
        ws['A2'].font = sub_font
        ws['A2'].alignment = center

        ws.merge_cells('A3:G3')
        ws['A3'] = f'Period: {period_text}'
        ws['A3'].font = sub_font
        ws['A3'].alignment = center

        ws.merge_cells('A4:D4')
        ws['A4'] = f'Doctor: {doctor_name}'
        ws['A4'].font = Font(bold=True, color='1F3A3D')
        ws['A4'].alignment = Alignment(horizontal='left')

        ws.merge_cells('E4:G4')
        ws['E4'] = f'Hospital: {hospital_name}'
        ws['E4'].font = Font(bold=True, color='1F3A3D')
        ws['E4'].alignment = Alignment(horizontal='right')

        # Header row
        headers = ['No.', 'Date', 'Description', 'Units', 'Work Type', 'Patient', 'Amount']
        header_row = 6
        for col_idx, h in enumerate(headers, start=1):
            c = ws.cell(row=header_row, column=col_idx, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = center
            c.border = border

        # Data rows
        total_amount = 0.0
        row = header_row + 1
        for idx, e in enumerate(entries, start=1):
            amount = float(e.amount or 0)
            total_amount += amount

            row_values = [
                idx,
                _format_entry_period(e),
                e.description or '',
                e.no_of_units or '',
                e.work_type or '',
                e.patient_name or '',
                amount,
            ]
            for col_idx, val in enumerate(row_values, start=1):
                c = ws.cell(row=row, column=col_idx, value=val)
                c.border = border
                if col_idx == 7:
                    c.number_format = '₹#,##0.00'
                    c.alignment = right
                elif col_idx in (1, 2, 4):
                    c.alignment = center
            row += 1

        # Totals row
        ws.cell(row=row, column=6, value='TOTAL AMOUNT')
        ws.cell(row=row, column=7, value=total_amount)
        for col_idx in range(1, 8):
            c = ws.cell(row=row, column=col_idx)
            c.font = Font(bold=True)
            c.fill = total_fill
            c.border = border
            if col_idx == 7:
                c.number_format = '₹#,##0.00'
                c.alignment = right
            elif col_idx == 6:
                c.alignment = right

        # Carry / Remaining rows
        row += 2
        ws.cell(row=row, column=6, value='Previous Balance (₹):').font = Font(bold=True)
        ws.cell(row=row, column=7, value=prev_balance).number_format = '₹#,##0.00'

        row += 1
        remaining = total_amount - prev_balance
        ws.cell(row=row, column=6, value='Remaining Amount (₹):').font = Font(bold=True)
        ws.cell(row=row, column=7, value=remaining).number_format = '₹#,##0.00'
        ws.cell(row=row, column=7).font = Font(bold=True)

        # Column widths
        widths = [6, 14, 42, 8, 16, 22, 14]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

        # Stream
        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)

        filename = f"Bill_{doctor_name or 'Doctor'}_{filename_suffix}.xlsx"
        return send_file(
            stream,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=filename,
        )

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': f'Export failed: {str(e)}'}), 500